"""Excel storage engine.

Every table is stored in its own workbook under backend/excel/<table>.xlsx
with a single worksheet named ``data``. Row 1 always holds the column headers
defined in utils.constants.TABLE_SCHEMAS. All access is guarded by a
per-table re-entrant lock so the API can be used from multiple threads.
"""
import os
import tempfile
import threading

from openpyxl import Workbook, load_workbook

from ..utils import constants as C
from ..utils.config import DEFAULT_ADMIN_PASSWORD, EXCEL_DIR
from ..utils.helpers import make_id, now_str, today_str

SHEET_NAME = "data"
_locks = {}
_global_lock = threading.RLock()
_cache = {}


def _lock_for(table):
    with _global_lock:
        if table not in _locks:
            _locks[table] = threading.RLock()
        return _locks[table]


def _path(table):
    return os.path.join(EXCEL_DIR, f"{table}.xlsx")


def _headers(table):
    return list(C.TABLE_SCHEMAS.get(table, []))


def _to_value(value):
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return value


def _row_to_dict(headers, values):
    row = {}
    for index, key in enumerate(headers):
        value = values[index] if index < len(values) else ""
        row[key] = "" if value is None else str(value).strip() if not isinstance(value, (int, float)) else value
    return row


def _atomic_save(wb, path):
    """Save to a temp file then atomically replace the target.

    Guarantees the workbook on disk is never observed in a torn state, even
    if another process writes the same table at the same time.
    """
    directory = os.path.dirname(path)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=directory)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except Exception:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise


def _ensure_workbook(table, wb=None):
    """Create the workbook with headers if missing; return a loaded workbook."""
    path = _path(table)
    headers = _headers(table)
    if not os.path.exists(path):
        os.makedirs(EXCEL_DIR, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(headers)
        _atomic_save(wb, path)
        _cache.pop(table, None)
        return load_workbook(path)
    wb = load_workbook(path)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    ws.title = SHEET_NAME
    current = [str(c.value).strip() if c.value else "" for c in ws[1]]
    if current != headers:
        # Schema changed: remap existing rows onto the new header order.
        existing = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
        wb.close()
        rows = []
        for values in existing:
            row = {}
            for idx, name in enumerate(current):
                if name:
                    row[name] = values[idx] if idx < len(values) else ""
            rows.append(row)
        write_all(table, rows)
        return load_workbook(path)
    return wb


def read_all(table, use_cache=False):
    """Return every row of ``table`` as a list of dicts."""
    if use_cache and table in _cache:
        return [dict(r) for r in _cache[table]]
    with _lock_for(table):
        wb = _ensure_workbook(table)
        ws = wb[SHEET_NAME]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        wb.close()
        headers = _headers(table)
        result = [_row_to_dict(headers, r) for r in rows if any(v not in (None, "") for v in r)]
        _cache[table] = result
        return [dict(r) for r in result]


def write_all(table, rows):
    """Replace the whole sheet content with ``rows``."""
    with _lock_for(table):
        headers = _headers(table)
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(headers)
        for row in rows:
            ws.append([_to_value(row.get(h, "")) for h in headers])
        os.makedirs(EXCEL_DIR, exist_ok=True)
        _atomic_save(wb, _path(table))
        _cache[table] = [dict(r) for r in rows]


def insert(table, row):
    """Append one row (missing schema columns are blank-filled). Returns the row."""
    headers = _headers(table)
    record = {h: _to_value(row.get(h, "")) for h in headers}
    with _lock_for(table):
        rows = read_all(table)
        rows.append(record)
        write_all(table, rows)
    return dict(record)


def find_all(table, **conditions):
    rows = read_all(table, use_cache=True)
    for key, value in conditions.items():
        if value in (None, ""):
            continue
        rows = [r for r in rows if str(r.get(key, "")).lower() == str(value).lower()]
    return rows


def find_one(table, **conditions):
    rows = find_all(table, **conditions)
    return rows[0] if rows else None


def find_by_id(table, row_id):
    if row_id in (None, ""):
        return None
    return find_one(table, id=str(row_id))


def update(table, row_id, changes):
    """Merge ``changes`` into the row with ``row_id``. Returns the row or None."""
    headers = _headers(table)
    with _lock_for(table):
        rows = read_all(table)
        target = None
        for row in rows:
            if str(row.get("id")) == str(row_id):
                target = row
                break
        if target is None:
            return None
        for key, value in changes.items():
            if key in headers and value is not None:
                target[key] = _to_value(value)
        target["updated_at"] = now_str()
        write_all(table, rows)
        return dict(target)


def delete(table, row_id):
    with _lock_for(table):
        rows = read_all(table)
        kept = [r for r in rows if str(r.get("id")) != str(row_id)]
        if len(kept) == len(rows):
            return False
        write_all(table, kept)
        return True


def count(table, **conditions):
    return len(find_all(table, **conditions))


def exists(table, **conditions):
    return find_one(table, **conditions) is not None


def next_numeric_id(rows):
    highest = 0
    for row in rows:
        try:
            highest = max(highest, int(row.get("id") or 0))
        except (TypeError, ValueError):
            continue
    return highest + 1


def seed_defaults():
    """Create all workbooks and seed the initial admin + role registry."""
    from werkzeug.security import generate_password_hash

    from ..utils.config import DEFAULT_ADMIN_EMAIL

    # Presets are only seeded on the very first run so a role the admin
    # deleted does not come back after a restart.
    fresh_install = not os.path.exists(_path("roles"))

    for table in C.TABLE_SCHEMAS:
        with _lock_for(table):
            _ensure_workbook(table)
        read_all(table)

    for role_name, description in (
        (C.ROLE_ADMIN, "Full access to every module, report and setting"),
        (C.ROLE_EMPLOYEE, "Sales and delivery staff: leads, companies, projects"),
        (C.ROLE_DEVELOPER, "Engineering team: assigned tasks and projects"),
    ):
        if not find_one("roles", name=role_name):
            insert("roles", {
                "id": make_id("role_"),
                "name": role_name,
                "access": role_name,
                "description": description,
                "built_in": "Yes",
                "created_at": now_str(),
                "updated_at": now_str(),
            })

    if fresh_install:
        for role_name, access, description in C.ROLE_PRESETS:
            if not find_one("roles", name=role_name):
                insert("roles", {
                    "id": make_id("role_"),
                    "name": role_name,
                    "access": access,
                    "description": description,
                    "built_in": "No",
                    "created_at": now_str(),
                    "updated_at": now_str(),
                })

    if not find_one("employees", email=DEFAULT_ADMIN_EMAIL):
        insert("employees", {
            "id": make_id("emp_"),
            "name": "System Admin",
            "email": DEFAULT_ADMIN_EMAIL,
            "password": generate_password_hash(DEFAULT_ADMIN_PASSWORD),
            "role": C.ROLE_ADMIN,
            "status": "Active",
            "department": "Administration",
            "designation": "Administrator",
            "joined_date": today_str(),
            "created_at": now_str(),
            "updated_at": now_str(),
        })
